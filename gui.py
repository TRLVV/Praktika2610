import tkinter as tk
from tkinter import ttk, messagebox, simpledialog
import psycopg2
from psycopg2 import IntegrityError
import bcrypt
import random
from datetime import datetime

# --- КОНФИГУРАЦИЯ БД ---
DB_CONFIG = {
    "host": "localhost",
    "port": 5432,
    "dbname": "postgres",
    "user": "user1",
    "password": "12345678"
}

# --- ЦВЕТА ТЁМНОЙ ТЕМЫ ---
BG_COLOR = "#1e1e1e"
FG_COLOR = "#e0e0e0"
BTN_BG = "#2d2d2d"
BTN_FG = "#ffffff"
ENTRY_BG = "#2d2d2d"
ENTRY_FG = "#e0e0e0"
ACCENT_COLOR = "#007acc"


def get_connection():
    try:
        return psycopg2.connect(**DB_CONFIG)
    except psycopg2.Error as e:
        messagebox.showerror("Ошибка БД", f"Не удалось подключиться:\n{e}")
        return None


def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')


def check_password(password: str, hashed: str) -> bool:
    return bcrypt.checkpw(password.encode('utf-8'), hashed.encode('utf-8'))


class App:
    def __init__(self, root):
        self.root = root
        self.root.title("Game System")
        self.root.geometry("900x600")
        self.root.configure(bg=BG_COLOR)

        # Настройка ttk стилей
        style = ttk.Style()
        style.theme_use('clam')
        style.configure('.', background=BG_COLOR, foreground=FG_COLOR, fieldbackground=ENTRY_BG)
        style.configure('TButton', background=BTN_BG, foreground=BTN_FG)
        style.configure('TLabel', background=BG_COLOR, foreground=FG_COLOR)
        style.configure('TEntry', fieldbackground=ENTRY_BG, foreground=ENTRY_FG)
        style.configure('Treeview', background=ENTRY_BG, foreground=FG_COLOR, fieldbackground=ENTRY_BG)
        style.configure('Treeview.Heading', background=BTN_BG, foreground=FG_COLOR)
        style.map('TButton', background=[('active', ACCENT_COLOR)])
        style.map('Treeview', background=[('selected', ACCENT_COLOR)], foreground=[('selected', '#ffffff')])

        self.conn = get_connection()
        if not self.conn:
            root.destroy()
            return

        self.current_user = None
        self.show_auth_frame()

    def clear_frame(self):
        for widget in self.root.winfo_children():
            widget.destroy()

    def create_button(self, parent, text, command, width=20):
        btn = tk.Button(parent, text=text, command=command, width=width,
                        bg=BTN_BG, fg=BTN_FG, activebackground=ACCENT_COLOR, activeforeground='#ffffff',
                        relief=tk.FLAT, padx=10, pady=5)
        return btn

    def create_label(self, parent, text, font=("Arial", 12), **kwargs):
        lbl = tk.Label(parent, text=text, font=font, bg=BG_COLOR, fg=FG_COLOR, **kwargs)
        return lbl

    def create_entry(self, parent, width=30, show=""):
        entry = tk.Entry(parent, width=width, bg=ENTRY_BG, fg=ENTRY_FG, insertbackground=FG_COLOR, show=show)
        return entry

    # --- ЭКСПОРТ В EXCEL ---
    def export_to_excel(self, data, filename, headers):
        if not data:
            messagebox.showwarning("Предупреждение", "Нет данных для экспорта.")
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = Workbook()
            ws = wb.active
            ws.title = "Data"

            # Стили
            header_font = Font(bold=True, color="FFFFFF", size=11)
            header_fill = PatternFill(start_color="#007acc", end_color="#007acc", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            thin_border = Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin')
            )

            # Заголовки
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            # Данные
            for row_num, row_data in enumerate(data, 2):
                for col_num, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border
                    if isinstance(value, datetime):
                        cell.number_format = 'YYYY-MM-DD HH:MM:SS'

            # Автоширина
            for col_cells in ws.columns:
                max_length = 0
                col_letter = col_cells[0].column_letter
                for cell in col_cells:
                    try:
                        length = len(str(cell.value or ""))
                        if length > max_length:
                            max_length = length
                    except:
                        pass
                adjusted_width = min(max_length + 4, 50)
                ws.column_dimensions[col_letter].width = adjusted_width

            filepath = f"{filename}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            wb.save(filepath)
            messagebox.showinfo("Успех", f"Файл сохранён:\n{filepath}")

        except ImportError:
            messagebox.showerror("Ошибка", "Библиотека openpyxl не установлена.\npip install openpyxl")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить файл:\n{e}")

    # --- АВТОРИЗАЦИЯ И РЕГИСТРАЦИЯ ---
    def show_auth_frame(self):
        self.clear_frame()
        self.current_user = None

        frame = tk.Frame(self.root, bg=BG_COLOR, padx=20, pady=20)
        frame.pack(expand=True)

        self.create_label(frame, "=== ГЛАВНОЕ МЕНЮ ===", font=("Arial", 16, "bold")).pack(pady=10)
        self.create_button(frame, "Войти", self.show_login).pack(pady=5)
        self.create_button(frame, "Регистрация", self.show_register).pack(pady=5)
        self.create_button(frame, "Выход", self.root.quit).pack(pady=5)

    def show_login(self):
        self.clear_frame()
        frame = tk.Frame(self.root, bg=BG_COLOR, padx=20, pady=20)
        frame.pack(expand=True)

        self.create_label(frame, "Вход", font=("Arial", 14, "bold")).pack(pady=10)

        self.create_label(frame, "Никнейм:").pack()
        username_entry = self.create_entry(frame)
        username_entry.pack(pady=5)

        self.create_label(frame, "Пароль:").pack()
        password_entry = self.create_entry(frame, show="*")
        password_entry.pack(pady=5)

        def do_login():
            username = username_entry.get().strip()
            password = password_entry.get()

            if not username or not password:
                messagebox.showwarning("Предупреждение", "Заполните все поля.")
                return

            cur = self.conn.cursor()
            try:
                cur.execute("""
                    SELECT u.user_id, u.username, u.password, r.role_name
                    FROM Users u
                    JOIN Roles r ON u.role_id = r.role_id
                    WHERE u.username = %s
                """, (username,))
                row = cur.fetchone()

                if not row or not check_password(password, row[2]):
                    messagebox.showerror("Ошибка", "Неверный никнейм или пароль.")
                    return

                user_id, uname, _, role_name = row

                # Проверка бана
                cur.execute("""
                    SELECT reason, ban_expires FROM Bans
                    WHERE banned_user_id = %s
                      AND (ban_expires IS NULL OR ban_expires > NOW())
                    ORDER BY banned_at DESC LIMIT 1
                """, (user_id,))
                ban = cur.fetchone()
                if ban:
                    expires = f"до {ban[1].strftime('%Y-%m-%d %H:%M')}" if ban[1] else "НАВСЕГДА"
                    messagebox.showerror("ЗАБАНЕН", f"Вы забанены {expires}\nПричина: {ban[0]}")
                    return

                # Создание сессии
                cur.execute(
                    "INSERT INTO Sessions (user_id) VALUES (%s) RETURNING session_id",
                    (user_id,)
                )
                session_id = cur.fetchone()[0]
                self.conn.commit()

                self.current_user = {
                    "user_id": user_id,
                    "username": uname,
                    "role": role_name,
                    "session_id": session_id
                }
                messagebox.showinfo("Успех", f"Добро пожаловать, {uname}!")
                self.show_main_menu()

            except psycopg2.Error as e:
                self.conn.rollback()
                messagebox.showerror("Ошибка", f"Ошибка входа:\n{e}")
            finally:
                cur.close()

        self.create_button(frame, "Войти", do_login).pack(pady=10)
        self.create_button(frame, "Назад", self.show_auth_frame).pack(pady=5)

    def show_register(self):
        self.clear_frame()
        frame = tk.Frame(self.root, bg=BG_COLOR, padx=20, pady=20)
        frame.pack(expand=True)

        self.create_label(frame, "Регистрация", font=("Arial", 14, "bold")).pack(pady=10)

        self.create_label(frame, "Никнейм:").pack()
        username_entry = self.create_entry(frame)
        username_entry.pack(pady=5)

        self.create_label(frame, "Пароль (мин. 6 символов):").pack()
        password_entry = self.create_entry(frame, show="*")
        password_entry.pack(pady=5)

        def do_register():
            username = username_entry.get().strip()
            password = password_entry.get()

            if not username:
                messagebox.showwarning("Предупреждение", "Никнейм не может быть пустым.")
                return
            if len(password) < 6:
                messagebox.showwarning("Предупреждение", "Пароль должен быть минимум 6 символов.")
                return

            cur = self.conn.cursor()
            try:
                cur.execute("SELECT role_id FROM Roles WHERE role_name = 'user'")
                role_row = cur.fetchone()
                if not role_row:
                    messagebox.showerror("Ошибка", "Роль 'user' не найдена в БД.")
                    return

                cur.execute("SELECT region_id FROM Regions ORDER BY region_id LIMIT 1")
                region_row = cur.fetchone()
                if not region_row:
                    messagebox.showerror("Ошибка", "Нет доступных регионов в БД.")
                    return

                cur.execute(
                    "INSERT INTO Users (username, password, role_id, region_id) VALUES (%s, %s, %s, %s)",
                    (username, hash_password(password), role_row[0], region_row[0])
                )
                self.conn.commit()
                messagebox.showinfo("Успех", f"Пользователь '{username}' зарегистрирован.")
                self.show_auth_frame()

            except IntegrityError:
                self.conn.rollback()
                messagebox.showerror("Ошибка", f"Никнейм '{username}' уже занят.")
            except psycopg2.Error as e:
                self.conn.rollback()
                messagebox.showerror("Ошибка", f"Ошибка регистрации:\n{e}")
            finally:
                cur.close()

        self.create_button(frame, "Зарегистрироваться", do_register).pack(pady=10)
        self.create_button(frame, "Назад", self.show_auth_frame).pack(pady=5)

    # --- ГЛАВНОЕ МЕНЮ ---
    def show_main_menu(self):
        self.clear_frame()
        frame = tk.Frame(self.root, bg=BG_COLOR, padx=20, pady=20)
        frame.pack(expand=True)

        is_admin = self.current_user["role"] == "admin"
        self.create_label(frame, f"МЕНЮ ({self.current_user['username']} | {self.current_user['role']})",
                          font=("Arial", 14, "bold")).pack(pady=10)

        self.create_button(frame, "Просмотр сессий", self.show_sessions, width=30).pack(pady=5)
        self.create_button(frame, "Результаты матчей", self.show_match_results, width=30).pack(pady=5)

        # Игры теперь обновляют текущее окно
        self.create_button(frame, "Играть в Кликер", self.play_clicker, width=30).pack(pady=5)
        self.create_button(frame, "Играть в Угадай число", self.play_guessnumber, width=30).pack(pady=5)
        self.create_button(frame, "Играть в Крестики-нолики", self.play_tictactoe, width=30).pack(pady=5)

        if is_admin:
            self.create_button(frame, "Забанить пользователя", self.admin_ban, width=30).pack(pady=5)
            self.create_button(frame, "Выдать роль мода", self.admin_promote, width=30).pack(pady=5)

        self.create_button(frame, "Выйти из аккаунта", self.logout, width=30).pack(pady=10)
        self.create_button(frame, "Выход из программы", self.root.quit, width=30).pack(pady=5)

    def logout(self):
        cur = self.conn.cursor()
        try:
            cur.execute(
                "UPDATE Sessions SET ended_at = NOW() WHERE session_id = %s AND ended_at IS NULL",
                (self.current_user["session_id"],)
            )
            self.conn.commit()
        except psycopg2.Error as e:
            print(f"Ошибка выхода: {e}")
        finally:
            cur.close()
        self.show_auth_frame()

    # --- ПРОСМОТР СЕССИЙ ---
    def show_sessions(self):
        self.clear_frame()
        frame = tk.Frame(self.root, bg=BG_COLOR, padx=10, pady=10)
        frame.pack(fill=tk.BOTH, expand=True)

        # 1. Сначала определяем структуру данных
        columns = ("ID", "Ник", "Игра", "Начало", "Конец")

        filter_frame = tk.Frame(frame, bg=BG_COLOR)
        filter_frame.pack(fill=tk.X, pady=5)

        self.create_label(filter_frame, "Фильтр по игре:").pack(side=tk.LEFT)
        game_var = tk.StringVar(value="Все")
        game_combo = ttk.Combobox(filter_frame, textvariable=game_var,
                                  values=["Все", "clicker", "tictactoe", "guessnumber"], width=15)
        game_combo.pack(side=tk.LEFT, padx=5)

        self.create_label(filter_frame, "Поиск по нику:").pack(side=tk.LEFT, padx=(10, 0))
        search_entry = self.create_entry(filter_frame, width=20)
        search_entry.pack(side=tk.LEFT, padx=5)

        # Создаем Treeview ДО кнопок, чтобы лямбды его видели
        tree = ttk.Treeview(frame, columns=columns, show="headings", height=20)
        for col in columns:
            tree.heading(col, text=col)
            tree.column(col, width=150)
        tree.pack(fill=tk.BOTH, expand=True)

        scrollbar = ttk.Scrollbar(frame, orient=tk.VERTICAL, command=tree.yview)
        tree.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        # Теперь определяем load_data, когда tree уже существует
        def load_data():
            for item in tree.get_children():
                tree.delete(item)

            game_filter = game_var.get().lower()
            search = search_entry.get().strip()

            conditions = []
            params = []

            if game_filter != "все":
                conditions.append("g.game_name = %s")
                params.append(game_filter)

            if search:
                conditions.append("u.username ILIKE %s")
                params.append(f"%{search}%")

            where_clause = "WHERE " + " AND ".join(conditions) if conditions else ""

            query = f"""
                SELECT s.session_id, u.username, g.game_name, s.started_at, s.ended_at
                FROM Sessions s
                JOIN Users u ON s.user_id = u.user_id
                LEFT JOIN Matches m ON m.session_id = s.session_id
                LEFT JOIN Games g ON m.game_id = g.game_id
                {where_clause}
                ORDER BY s.started_at DESC
                LIMIT 100;
            """

            cur = self.conn.cursor()
            try:
                cur.execute(query, params)
                rows = cur.fetchall()
                for r in rows:
                    ended = r[4].strftime('%Y-%m-%d %H:%M:%S') if r[4] else "Активна"
                    started = r[3].strftime('%Y-%m-%d %H:%M:%S')
                    game = r[2] or "-"
                    tree.insert("", tk.END, values=(r[0], r[1], game, started, ended))
            except psycopg2.Error as e:
                messagebox.showerror("Ошибка", f"{e}")
            finally:
                cur.close()

        # 3. И только ТЕПЕРЬ создаем кнопки, ссылаясь на уже существующие tree и load_data
        self.create_button(filter_frame, "Экспорт в Excel",
                           lambda: self.export_to_excel(
                               [tree.item(item)['values'] for item in tree.get_children()],
                               "sessions",
                               columns
                           )).pack(side=tk.LEFT, padx=5)

        self.create_button(filter_frame, "Применить фильтр", load_data).pack(side=tk.LEFT, padx=5)

        load_data()  # Загружаем данные сразу

        self.create_button(frame, "Назад", self.show_main_menu).pack(pady=10)

    # --- РЕЗУЛЬТАТЫ МАТЧЕЙ ---
    def show_match_results(self):
        self.clear_frame()
        frame = tk.Frame(self.root, bg=BG_COLOR, padx=10, pady=10)
        frame.pack(fill=tk.BOTH, expand=True)

        # 1. Структура
        columns = ("MatchID", "Ник", "Игра", "Результат", "Клики", "Попытки", "Дата")

        filter_frame = tk.Frame(frame, bg=BG_COLOR)
        filter_frame.pack(fill=tk.X, pady=5)

        self.create_label(filter_frame, "Фильтр по игре:").pack(side=tk.LEFT)
        game_var = tk.StringVar(value="Все")
        game_combo = ttk.Combobox(filter_frame, textvariable=game_var,
                                  values=["Все", "clicker", "tictactoe", "guessnumber"], width=15)
        game_combo.pack(side=tk.LEFT, padx=5)

        self.create_label(filter_frame, "Поиск по нику:").pack(side=tk.LEFT, padx=(10, 0))
        search_entry = self.create_entry(filter_frame, width=20)
        search_entry.pack(side=tk.LEFT, padx=5)

        # 2. Treeview создаем ДО кнопок
        tree = ttk.Treeview(frame, columns=columns, show="headings", height=20)
        for col in columns:
            tree.heading(col, text=col)
            tree.column(col, width=120)
        tree.pack(fill=tk.BOTH, expand=True)

        scrollbar = ttk.Scrollbar(frame, orient=tk.VERTICAL, command=tree.yview)
        tree.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        # 3. Функция загрузки
        def load_data():
            for item in tree.get_children():
                tree.delete(item)

            game_filter = game_var.get().lower()
            search = search_entry.get().strip()

            conditions = []
            params = []

            if game_filter != "все":
                conditions.append("g.game_name = %s")
                params.append(game_filter)

            if search:
                conditions.append("u.username ILIKE %s")
                params.append(f"%{search}%")

            where_clause = "WHERE " + " AND ".join(conditions) if conditions else ""

            query = f"""
                SELECT ms.match_id, u.username, g.game_name, ms.result,
                       ms.click_count, ms.guess_attempts, m.started_at
                FROM MatchStats ms
                JOIN Users u ON ms.user_id = u.user_id
                JOIN Matches m ON ms.match_id = m.match_id
                JOIN Games g ON m.game_id = g.game_id
                {where_clause}
                ORDER BY m.started_at DESC
                LIMIT 100;
            """

            cur = self.conn.cursor()
            try:
                cur.execute(query, params)
                rows = cur.fetchall()
                for r in rows:
                    clicks = str(r[4]) if r[4] is not None else "-"
                    guesses = str(r[5]) if r[5] is not None else "-"
                    result = r[3] if r[3] is not None else "-"
                    date = r[6].strftime('%Y-%m-%d %H:%M')
                    tree.insert("", tk.END, values=(r[0], r[1], r[2], result, clicks, guesses, date))
            except psycopg2.Error as e:
                messagebox.showerror("Ошибка", f"{e}")
            finally:
                cur.close()

        # 4. Кнопки (теперь tree и load_data существуют)
        self.create_button(filter_frame, "Экспорт в Excel",
                           lambda: self.export_to_excel(
                               [tree.item(item)['values'] for item in tree.get_children()],
                               "match_stats",
                               columns
                           )).pack(side=tk.LEFT, padx=5)

        self.create_button(filter_frame, "Применить фильтр", load_data).pack(side=tk.LEFT, padx=5)

        load_data()

        self.create_button(frame, "Назад", self.show_main_menu).pack(pady=10)

    # --- ЛОГИКА ИГР (ОБЩИЕ МЕТОДЫ) ---
    def create_match(self, game_name):
        cur = self.conn.cursor()
        try:
            cur.execute("SELECT game_id FROM Games WHERE game_name = %s", (game_name,))
            row = cur.fetchone()
            if not row:
                messagebox.showerror("Ошибка", f"Игра '{game_name}' не найдена в БД.")
                return None

            cur.execute(
                "INSERT INTO Matches (session_id, game_id) VALUES (%s, %s) RETURNING match_id",
                (self.current_user["session_id"], row[0])
            )
            match_id = cur.fetchone()[0]
            self.conn.commit()
            return match_id
        except psycopg2.Error as e:
            self.conn.rollback()
            messagebox.showerror("Ошибка", f"Не удалось создать матч:\n{e}")
            return None
        finally:
            cur.close()

    def save_match_result(self, match_id, user_id, click_count=None, guess_attempts=None, result=None):
        cur = self.conn.cursor()
        try:
            cur.execute(
                """INSERT INTO MatchStats (match_id, user_id, click_count, guess_attempts, result)
                   VALUES (%s, %s, %s, %s, %s)""",
                (match_id, user_id, click_count, guess_attempts, result)
            )
            cur.execute(
                "UPDATE Matches SET ended_at = NOW() WHERE match_id = %s",
                (match_id,)
            )
            self.conn.commit()
        except psycopg2.Error as e:
            self.conn.rollback()
            messagebox.showerror("Ошибка", f"Не удалось сохранить результат:\n{e}")
        finally:
            cur.close()

    # --- ИГРА: КЛИКЕР ---
    def play_clicker(self):
        self.clear_frame()
        frame = tk.Frame(self.root, bg=BG_COLOR, padx=20, pady=20)
        frame.pack(expand=True)

        match_id = self.create_match('clicker')
        if not match_id:
            self.show_main_menu()
            return

        clicks = 0
        time_left = 10
        running = False

        label = self.create_label(frame, "Нажмите Старт", font=("Arial", 20))
        label.pack(pady=20)

        clicks_label = self.create_label(frame, f"Кликов: {clicks}", font=("Arial", 16))
        clicks_label.pack(pady=10)

        time_label = self.create_label(frame, f"Время: {time_left}с", font=("Arial", 16))
        time_label.pack(pady=10)

        def on_click():
            nonlocal clicks
            if running:
                clicks += 1
                clicks_label.config(text=f"Кликов: {clicks}")

        def countdown():
            nonlocal time_left, running
            if time_left > 0 and running:
                time_left -= 1
                time_label.config(text=f"Время: {time_left}с")
                self.root.after(1000, countdown)
            elif running:
                running = False
                self.save_match_result(match_id, self.current_user["user_id"], click_count=clicks)
                label.config(text=f"Готово! Кликов: {clicks}")
                click_btn.config(state=tk.DISABLED)
                start_btn.config(state=tk.DISABLED)
                messagebox.showinfo("Результат", f"Вы сделали {clicks} кликов!")
                # Возврат в меню после закрытия сообщения
                self.root.after(100, self.show_main_menu)

        def start_game():
            nonlocal running, time_left, clicks
            running = True
            time_left = 10
            clicks = 0
            clicks_label.config(text=f"Кликов: {clicks}")
            time_label.config(text=f"Время: {time_left}с")
            label.config(text="Кликайте!")
            start_btn.config(state=tk.DISABLED)
            countdown()

        click_btn = tk.Button(frame, text="КЛИК!", width=20, height=3, command=on_click,
                              bg=BTN_BG, fg=BTN_FG, activebackground=ACCENT_COLOR, relief=tk.FLAT)
        click_btn.pack(pady=10)

        start_btn = tk.Button(frame, text="Старт", width=20, command=start_game,
                              bg=BTN_BG, fg=BTN_FG, activebackground=ACCENT_COLOR, relief=tk.FLAT)
        start_btn.pack(pady=10)

        self.create_button(frame, "Отмена", self.show_main_menu).pack(pady=10)

    # --- ИГРА: УГАДАЙ ЧИСЛО ---
    def play_guessnumber(self):
        self.clear_frame()
        frame = tk.Frame(self.root, bg=BG_COLOR, padx=20, pady=20)
        frame.pack(expand=True)

        match_id = self.create_match('guessnumber')
        if not match_id:
            self.show_main_menu()
            return

        secret = random.randint(1, 100)
        attempts = 0
        max_attempts = 10

        label = self.create_label(frame, f"Я загадал число от 1 до 100.\nУ тебя {max_attempts} попыток.",
                                  font=("Arial", 14))
        label.pack(pady=20)

        entry = self.create_entry(frame, width=10)
        entry.config(font=("Arial", 16))
        entry.pack(pady=10)

        result_label = self.create_label(frame, "", font=("Arial", 14))
        result_label.pack(pady=10)

        def check_guess():
            nonlocal attempts
            try:
                guess = int(entry.get())
            except ValueError:
                messagebox.showwarning("Ошибка", "Введите число.")
                return

            attempts += 1
            if guess == secret:
                self.save_match_result(match_id, self.current_user["user_id"], guess_attempts=attempts, result='win')
                result_label.config(text=f"Угадал за {attempts} попыток!")
                entry.config(state=tk.DISABLED)
                btn.config(state=tk.DISABLED)
                messagebox.showinfo("Победа", f"Число было {secret}. Попыток: {attempts}")
                self.root.after(100, self.show_main_menu)
            elif attempts >= max_attempts:
                self.save_match_result(match_id, self.current_user["user_id"], guess_attempts=attempts, result='loss')
                result_label.config(text=f"Проиграл. Число было {secret}.")
                entry.config(state=tk.DISABLED)
                btn.config(state=tk.DISABLED)
                messagebox.showinfo("Проигрыш", f"Число было {secret}.")
                self.root.after(100, self.show_main_menu)
            elif guess < secret:
                result_label.config(text=f"Больше. Попытка {attempts}/{max_attempts}")
            else:
                result_label.config(text=f"Меньше. Попытка {attempts}/{max_attempts}")

            entry.delete(0, tk.END)

        btn = tk.Button(frame, text="Проверить", width=20, command=check_guess,
                        bg=BTN_BG, fg=BTN_FG, activebackground=ACCENT_COLOR, relief=tk.FLAT)
        btn.pack(pady=10)

        self.create_button(frame, "Отмена", self.show_main_menu).pack(pady=10)

    # --- ИГРА: КРЕСТИКИ-НОЛИКИ ---
    def play_tictactoe(self):
        self.clear_frame()
        frame = tk.Frame(self.root, bg=BG_COLOR, padx=20, pady=20)
        frame.pack(expand=True)

        match_id = self.create_match('tictactoe')
        if not match_id:
            self.show_main_menu()
            return

        board = [' '] * 9
        buttons = []

        def check_win(b, mark):
            wins = [(0, 1, 2), (3, 4, 5), (6, 7, 8), (0, 3, 6), (1, 4, 7), (2, 5, 8), (0, 4, 8), (2, 4, 6)]
            return any(b[a] == b[b_] == b[c] == mark for a, b_, c in wins)

        def on_click(idx):
            if board[idx] != ' ' or buttons[idx]['state'] == tk.DISABLED:
                return

            board[idx] = 'X'
            buttons[idx].config(text='X', state=tk.DISABLED)

            if check_win(board, 'X'):
                self.save_match_result(match_id, self.current_user["user_id"], result='win')
                messagebox.showinfo("Победа", "Вы победили!")
                disable_all()
                self.root.after(100, self.show_main_menu)
                return

            if ' ' not in board:
                self.save_match_result(match_id, self.current_user["user_id"], result='draw')
                messagebox.showinfo("Ничья", "Ничья!")
                disable_all()
                self.root.after(100, self.show_main_menu)
                return

            empty = [i for i, x in enumerate(board) if x == ' ']
            bot_move = random.choice(empty)
            board[bot_move] = 'O'
            buttons[bot_move].config(text='O', state=tk.DISABLED)

            if check_win(board, 'O'):
                self.save_match_result(match_id, self.current_user["user_id"], result='loss')
                messagebox.showinfo("Проигрыш", "Компьютер победил.")
                disable_all()
                self.root.after(100, self.show_main_menu)
                return

            if ' ' not in board:
                self.save_match_result(match_id, self.current_user["user_id"], result='draw')
                messagebox.showinfo("Ничья", "Ничья!")
                disable_all()
                self.root.after(100, self.show_main_menu)
                return

        def disable_all():
            for b in buttons:
                b.config(state=tk.DISABLED)

        grid_frame = tk.Frame(frame, bg=BG_COLOR)
        grid_frame.pack(pady=20)

        for i in range(9):
            btn = tk.Button(grid_frame, text=' ', width=10, height=3, font=("Arial", 20),
                            command=lambda idx=i: on_click(idx),
                            bg=BTN_BG, fg=BTN_FG, activebackground=ACCENT_COLOR, relief=tk.FLAT)
            btn.grid(row=i // 3, column=i % 3, padx=5, pady=5)
            buttons.append(btn)

        self.create_button(frame, "Отмена", self.show_main_menu).pack(pady=10)

    # --- АДМИН ПАНЕЛЬ ---
    def admin_ban(self):
        target = simpledialog.askstring("Бан", "Никнейм пользователя для бана:")
        if not target:
            return

        reason = simpledialog.askstring("Бан", "Причина бана:")
        if not reason:
            messagebox.showwarning("Предупреждение", "Причина обязательна.")
            return

        cur = self.conn.cursor()
        try:
            cur.execute("SELECT user_id FROM Users WHERE username = %s", (target,))
            target_row = cur.fetchone()
            if not target_row:
                messagebox.showerror("Ошибка", f"Пользователь '{target}' не найден.")
                return

            target_id = target_row[0]
            if target_id == self.current_user["user_id"]:
                messagebox.showwarning("Предупреждение", "Нельзя забанить самого себя.")
                return

            cur.execute(
                "INSERT INTO Bans (banned_user_id, banned_by_user_id, reason) VALUES (%s, %s, %s)",
                (target_id, self.current_user["user_id"], reason)
            )
            self.conn.commit()
            messagebox.showinfo("Успех", f"Пользователь '{target}' забанен навсегда.")
        except IntegrityError as e:
            self.conn.rollback()
            messagebox.showerror("Ошибка", f"{e}")
        except psycopg2.Error as e:
            self.conn.rollback()
            messagebox.showerror("Ошибка", f"{e}")
        finally:
            cur.close()

    def admin_promote(self):
        target = simpledialog.askstring("Повышение", "Никнейм пользователя для выдачи мода:")
        if not target:
            return

        cur = self.conn.cursor()
        try:
            cur.execute("SELECT user_id FROM Users WHERE username = %s", (target,))
            target_row = cur.fetchone()
            if not target_row:
                messagebox.showerror("Ошибка", f"Пользователь '{target}' не найден.")
                return

            cur.execute("SELECT role_id FROM Roles WHERE role_name = 'mod'")
            mod_role = cur.fetchone()
            if not mod_role:
                messagebox.showerror("Ошибка", "Роль 'mod' не найдена в БД.")
                return

            cur.execute(
                "UPDATE Users SET role_id = %s WHERE user_id = %s",
                (mod_role[0], target_row[0])
            )
            self.conn.commit()
            if cur.rowcount == 0:
                messagebox.showwarning("Предупреждение", f"Пользователь '{target}' не обновлён.")
            else:
                messagebox.showinfo("Успех", f"Пользователю '{target}' выдана роль 'mod'.")
        except psycopg2.Error as e:
            self.conn.rollback()
            messagebox.showerror("Ошибка", f"{e}")
        finally:
            cur.close()


if __name__ == "__main__":
    root = tk.Tk()
    app = App(root)
    root.mainloop()
